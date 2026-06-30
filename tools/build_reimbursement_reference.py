from pathlib import Path
import subprocess

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.oxml.table import CT_Tbl
from docx.oxml.text.paragraph import CT_P
from docx.table import Table
from docx.text.paragraph import Paragraph
from PIL import Image, ImageDraw, ImageFont
from docx.shared import Pt, Cm
from pypdf import PdfReader, PdfWriter


ROOT = Path("/Users/hao/Documents/自助报销Agent")
DOWNLOADS = Path("/Users/hao/Downloads")
SOURCE_PDF = DOWNLOADS / "2018年新版常见业务费报销指南(文字版).pdf"
OUT_DOCX = ROOT / "output/docx/上海大学报销参考整理.docx"
OUT_NOTE_PDF = ROOT / "output/pdf/上海大学报销参考整理.pdf"
OUT_APPENDIX_PDF = ROOT / "output/pdf/2018年新版常见业务费报销指南(文字版)-已补充报销参考.pdf"


def set_run_font(run, size=10.5, bold=False):
    run.font.name = "Songti SC"
    run.font.size = Pt(size)
    run.bold = bold
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.rFonts
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    rfonts.set(qn("w:eastAsia"), "Songti SC")


def add_heading(doc, text, level=1):
    p = doc.add_paragraph()
    p.style = f"Heading {level}"
    run = p.add_run(text)
    set_run_font(run, 16 if level == 1 else 13, True)
    return p


def add_para(doc, text="", bold_prefix=None):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.line_spacing = 1.15
    if bold_prefix and text.startswith(bold_prefix):
        r1 = p.add_run(bold_prefix)
        set_run_font(r1, 10.5, True)
        r2 = p.add_run(text[len(bold_prefix):])
        set_run_font(r2)
    else:
        run = p.add_run(text)
        set_run_font(run)
    return p


def add_bullet(doc, text):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(2)
    run = p.add_run(text)
    set_run_font(run)


def add_number(doc, text):
    p = doc.add_paragraph(style="List Number")
    p.paragraph_format.space_after = Pt(2)
    run = p.add_run(text)
    set_run_font(run)


def add_table(doc, headers, rows):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    hdr = table.rows[0].cells
    for i, h in enumerate(headers):
        run = hdr[i].paragraphs[0].add_run(h)
        set_run_font(run, 10, True)
    for row in rows:
        cells = table.add_row().cells
        for i, value in enumerate(row):
            run = cells[i].paragraphs[0].add_run(value)
            set_run_font(run, 9.5)
    doc.add_paragraph()


def build_docx():
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Cm(1.8)
    section.bottom_margin = Cm(1.8)
    section.left_margin = Cm(2.0)
    section.right_margin = Cm(2.0)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("上海大学报销参考整理")
    set_run_font(run, 18, True)
    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = subtitle.add_run("用于上海大学自助报销 Agent 的资料准备清单和填报步骤")
    set_run_font(run, 11)

    add_para(doc, "整理日期：2026-06-30")
    add_para(doc, "适用边界：本整理用于准备材料、生成填报草稿和人工核对。涉及登录、审批、提交、借款或报销承诺等财务动作时，应由本人在上海大学财务系统内确认后操作。")

    add_heading(doc, "一、官方参考入口", 1)
    add_bullet(doc, "财务处首页：https://cwc.shu.edu.cn/")
    add_bullet(doc, "基本报销业务流程指南：https://cwc.shu.edu.cn/fwzn/jbbxywlczn.htm")
    add_bullet(doc, "系统操作指南：https://cwc.shu.edu.cn/fwzn/xtczzn.htm")
    add_bullet(doc, "2026-06-24 最新入口：智能报销系统操作演示PPT（202606），页面附件下载需验证码。")
    add_bullet(doc, "2025-11-05：2025年上海大学财务与资产业务百问百答，页面附件下载需验证码。")
    add_bullet(doc, "2025-03-31：上海大学智能报销操作手册（2025年3月），已下载并用于提取系统流程。")
    add_bullet(doc, "2018-10-10：2018年新版常见业务费报销指南(文字版)，即本次附加整理的原始 PDF。")

    add_heading(doc, "二、智能报销系统填报流程", 1)
    add_number(doc, "登录入口：手机端从企业微信进入工作台的智能财务；电脑端从学校主页进入信息门户、财务门户，再进入智能报销。")
    add_number(doc, "先处理票据：报销前把电子发票源文件、纸质发票照片和附件上传到发票夹。电子发票开具后通常需要等待约 24 小时才能完成查验。")
    add_number(doc, "选择业务单据：日常费用进入日常报销，国内差旅进入差旅报销，借款进入借款模块。")
    add_number(doc, "填写事由和预算项目：按真实业务说明用途，选择可用经费或预算项目。")
    add_number(doc, "勾选票据和附件：只有已查验、无需查验或系统允许使用的票据才能进入报销；支付记录、明细清单、通知、审批单、合同等按系统提示上传。")
    add_number(doc, "确认业务支出项：系统会按票据和项目匹配支出类型，仍需人工核对支出项是否对应真实业务。")
    add_number(doc, "填写支付信息：发票报销进卡通常带出本人银行卡；对公汇款需核对单位名称、开户行、账号等信息。")
    add_number(doc, "选择审批方式：默认线上审批；线下审批需预审后预约、打印、签字并投递纸质材料。")
    add_number(doc, "提交预审并跟踪：预审不通过时查看预审报告补材料；审批中可通过单据跟踪查看流转状态。")
    add_number(doc, "纸质票据规则：若报销单包含纸质票据，通常仍需打印报销单并把全部纸质材料投递；全电子材料按系统和财务处要求执行。")

    add_heading(doc, "三、报销材料速查表", 1)
    add_table(
        doc,
        ["业务类型", "核心材料", "特别注意"],
        [
            ["市内差旅", "交通结算单、出租车/公交/地铁等票据，必要时附审批或说明", "同车号出租车票通常不予报销；维持性经费开支需注意结算单要求。"],
            ["国内差旅", "国内出差审批单、差旅费结算单、往返交通票、住宿发票、会议通知及注册费发票", "老师、学生、校外人员材料不同；校外人员通常不领取伙食和交通补助。"],
            ["资料费/印刷费", "发票、图书或印刷明细、必要时合同", "图书和打印复印金额达到规定阈值时需附清单或合同。"],
            ["办公用品/耗材/电脑配件", "发票、品名单价数量明细", "发票内容需真实、完整；大额或笼统品类需补明细。"],
            ["版面费/出版费", "发票、录用通知、出版合同，外币支付时附外文票据、说明、录用通知和付款记录", "外币支付需先按财务处要求核算人民币后报销。"],
            ["测试化验加工费", "发票、对方单位盖章的测试/加工清单或合同", "清单或合同需能说明服务内容。"],
            ["会议费", "会议审批表、通知/议程/成果/人数/开支明细、签到表、支付记录", "一次会议原则上一次报销，不宜拆分多次提交。"],
            ["设备购置", "发票、固定资产审批单、订购合同、验收单，政府采购材料如适用", "固定资产、两用物资、家具类有不同登记和验收路径。"],
            ["数据采集费", "调查样本、问卷或案例材料、数据采集支出说明", "主要面向文科经费，说明需负责人签字。"],
            ["学生活动费", "原始票据、学生活动开支表", "通常仅维持性经费可支出，并需学生工作分管领导签字。"],
            ["公务接待费", "公务接待清单、餐费/食品/内转票据、支付记录、邀请函或来访公函", "维持性经费一事一报；横向科研经费单次达到阈值时需接待清单。"],
            ["国际合作交流费", "任务批件、预算审批意见、行程安排、机票行程单、护照出入境记录、签证保险票据、邀请信等", "国外票据需中文注明内容、日期、数量、金额并由经办人签字；一份任务批件通常一次核销。"],
            ["劳务费", "人员经费发放明细表、合同、专家咨询意见、讲座通知或海报等", "一份发放明细对应一份报销单，不能随意合并。"],
            ["借款业务", "借款联、销账联、合同/订单/立项书/通知/批文/审批单等附件", "取得发票后应及时按原借款校区办理冲销。"],
        ],
    )

    add_heading(doc, "四、通用合规检查", 1)
    checks = [
        "票据抬头应为上海大学，票据真实、合法、信息完整。",
        "发票内容应按实际业务开具，不能用办公用品、食品等笼统名称替代明细。",
        "公务卡、银行付款、POS 小票、银行对账单或短信截屏等支付记录应按业务要求留存。",
        "同一家公司同一天连号或相近票据不应拆分规避限额；达到限额时补清单、合同和支付记录。",
        "当年度发票一般应在次年 3 月底前处理，逾期风险需提前向财务处确认。",
        "每张报销单附件数量、纸质投递、线上审批、预约状态，以智能报销系统预审和财务处最新通知为准。",
    ]
    for item in checks:
        add_bullet(doc, item)

    add_heading(doc, "五、给上海大学自助报销 Agent 的字段清单", 1)
    add_table(
        doc,
        ["字段", "说明"],
        [
            ["报销人/经办人", "姓名、工号、联系方式、银行卡信息。"],
            ["业务类型", "日常报销、差旅报销、借款、劳务或其他专项业务。"],
            ["经费/预算项目", "项目号、负责人、可报支出项、余额和授权状态。"],
            ["事由", "业务发生原因、时间、地点、参与人、成果或用途。"],
            ["票据", "发票类型、号码、金额、开票日期、抬头、税号、查验状态、是否纸质。"],
            ["附件", "合同、通知、审批单、清单、支付记录、行程单、签到表、验收单等。"],
            ["支付方式", "进个人银行卡、对公汇款、内转、冲借款等。"],
            ["审批路径", "线上审批人、线下签字人、是否需要预约和投递。"],
            ["风险提示", "缺少材料、票据未查验、金额阈值触发清单/合同、纸质票据需投递等。"],
        ],
    )

    add_heading(doc, "六、自动报销可执行边界", 1)
    add_bullet(doc, "可以自动做：读取票据和附件、生成材料清单、提示缺漏、生成报销事由草稿、辅助在浏览器中逐项填写。")
    add_bullet(doc, "需要本人确认：登录、选择经费项目、确认票据真实性、确认支付信息、提交预审、提交审批、线下签字和投递。")
    add_bullet(doc, "不应自动做：保存或索要密码、绕过验证码、擅自提交财务单据、替本人作出真实性承诺。")

    add_heading(doc, "七、来源", 1)
    add_bullet(doc, "上海大学财务处首页，访问日期 2026-06-30：https://cwc.shu.edu.cn/")
    add_bullet(doc, "基本报销业务流程指南，访问日期 2026-06-30：https://cwc.shu.edu.cn/fwzn/jbbxywlczn.htm")
    add_bullet(doc, "系统操作指南，访问日期 2026-06-30：https://cwc.shu.edu.cn/fwzn/xtczzn.htm")
    add_bullet(doc, "上海大学智能报销操作手册（2025年3月），财务处附件 PDF。")
    add_bullet(doc, "2018年新版常见业务费报销指南(文字版)，财务处附件 PDF。")

    OUT_DOCX.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT_DOCX)


def convert_docx_to_pdf():
    OUT_NOTE_PDF.parent.mkdir(parents=True, exist_ok=True)
    doc = Document(OUT_DOCX)
    body = doc.element.body
    blocks = []
    list_no = 0
    for child in body.iterchildren():
        if isinstance(child, CT_P):
            para = Paragraph(child, doc)
            text = para.text.strip()
            if not text:
                continue
            style = para.style.name if para.style else ""
            if "List Bullet" in style:
                text = "• " + text
            elif "List Number" in style:
                list_no += 1
                text = f"{list_no}. {text}"
            else:
                list_no = 0
            blocks.append(("heading" if "Heading" in style else "text", text))
        elif isinstance(child, CT_Tbl):
            table = Table(child, doc)
            for row_idx, row in enumerate(table.rows):
                cells = [cell.text.strip().replace("\n", " / ") for cell in row.cells]
                blocks.append(("table_header" if row_idx == 0 else "table", " | ".join(cells)))
            blocks.append(("space", ""))

    font_path = "/System/Library/Fonts/STHeiti Light.ttc"
    font_regular = ImageFont.truetype(font_path, 24)
    font_heading = ImageFont.truetype("/System/Library/Fonts/STHeiti Medium.ttc", 34)
    font_table = ImageFont.truetype(font_path, 21)
    font_title = ImageFont.truetype("/System/Library/Fonts/STHeiti Medium.ttc", 42)

    page_w, page_h = 1240, 1754
    margin_x, margin_top, margin_bottom = 96, 92, 82
    line_gap = 10
    max_w = page_w - 2 * margin_x
    pages = []
    img = Image.new("RGB", (page_w, page_h), "white")
    draw = ImageDraw.Draw(img)
    y = margin_top

    def measure(text, font):
        return draw.textbbox((0, 0), text, font=font)[2]

    def wrap(text, font, width):
        lines, line = [], ""
        for ch in text:
            candidate = line + ch
            if measure(candidate, font) <= width or not line:
                line = candidate
            else:
                lines.append(line)
                line = ch
        if line:
            lines.append(line)
        return lines

    def new_page():
        nonlocal img, draw, y
        pages.append(img)
        img = Image.new("RGB", (page_w, page_h), "white")
        draw = ImageDraw.Draw(img)
        y = margin_top

    for kind, text in blocks:
        if kind == "space":
            y += 18
            continue
        font = font_regular
        extra_before = 0
        extra_after = 8
        if kind == "heading":
            font = font_heading
            extra_before = 18
            extra_after = 12
        elif kind == "table_header":
            font = font_table
            extra_before = 8
            extra_after = 7
        elif kind == "table":
            font = font_table
            extra_after = 5
        elif len(pages) == 0 and y == margin_top:
            font = font_title
            extra_after = 18

        wrapped = wrap(text, font, max_w)
        line_h = font.size + line_gap
        needed = extra_before + len(wrapped) * line_h + extra_after
        if y + needed > page_h - margin_bottom:
            new_page()
        y += extra_before
        for line in wrapped:
            draw.text((margin_x, y), line, fill=(20, 20, 20), font=font)
            y += line_h
        y += extra_after

    pages.append(img)
    pages[0].save(OUT_NOTE_PDF, "PDF", resolution=150.0, save_all=True, append_images=pages[1:])


def append_pdf():
    writer = PdfWriter()
    for pdf in (SOURCE_PDF, OUT_NOTE_PDF):
        reader = PdfReader(str(pdf))
        for page in reader.pages:
            writer.add_page(page)
    with OUT_APPENDIX_PDF.open("wb") as f:
        writer.write(f)


def main():
    build_docx()
    convert_docx_to_pdf()
    append_pdf()
    print(OUT_DOCX)
    print(OUT_NOTE_PDF)
    print(OUT_APPENDIX_PDF)
    print("pages", len(PdfReader(str(OUT_APPENDIX_PDF)).pages))


if __name__ == "__main__":
    main()
