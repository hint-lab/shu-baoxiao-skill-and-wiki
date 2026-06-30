#!/usr/bin/env python3
from __future__ import annotations

import argparse
import json
import re
import shutil
import sys
from collections import Counter
from dataclasses import asdict, dataclass
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

from pypdf import PdfReader

try:
    from docx import Document as DocxDocument
except Exception:  # pragma: no cover - optional dependency
    DocxDocument = None

try:
    from PIL import Image
    import pytesseract
except Exception:  # pragma: no cover - optional dependency
    Image = None
    pytesseract = None

ROOT = Path("/Users/hao/Documents/自助报销Agent")
RULES_PATH = ROOT / "data/rules/reimbursement_rules.json"
INDEX_PATH = ROOT / "data/cwc/index.json"
OUTPUT_ROOT = ROOT / "output/packages"

if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))
from tools.cwc_rag import bm25_scores, load_index, snippet  # noqa: E402
from tools.doubao_extractor import extract_fields_with_doubao  # noqa: E402


@dataclass
class SourceEvidence:
    title: str
    url: str
    source_type: str
    snippet: str
    score: float


@dataclass
class ImportedDocument:
    id: str
    file_name: str
    source_path: str
    document_type: str
    confidence: float
    text_preview: str
    extracted_fields: dict[str, Any]


@dataclass
class ChecklistItem:
    id: str
    name: str
    required: bool
    status: str
    matched_documents: list[str]
    why: str
    evidence: SourceEvidence | None


@dataclass
class ReimbursementPackage:
    id: str
    title: str
    created_at: str
    input_dir: str
    reimbursement_type: str
    reimbursement_type_name: str
    school_form_hint: str
    confidence: float
    user_note: str
    reason_draft: str
    documents: list[ImportedDocument]
    checklist: list[ChecklistItem]
    risk_flags: list[str]
    policy_evidence: list[SourceEvidence]


DOCUMENT_TYPE_KEYWORDS = {
    "invoice": [
        "发票",
        "invoice",
        "receipt",
        "发票号码",
        "开票日期",
        "价税合计",
        "购买方",
        "销售方",
        "电子发票",
        "财政票据",
    ],
    "payment_record": [
        "支付",
        "付款",
        "银行",
        "公务卡",
        "刷卡",
        "pos",
        "交易",
        "转账",
        "消费记录",
        "银行卡",
        "汇款",
    ],
    "approval": [
        "审批",
        "申请表",
        "出差审批",
        "预算审批",
        "固定资产",
        "验收",
        "登记",
    ],
    "settlement": [
        "结算单",
        "差旅费结算",
        "交通补助",
        "伙食补助",
    ],
    "itinerary": [
        "行程",
        "火车票",
        "高铁",
        "机票",
        "航空运输电子客票",
        "酒店",
        "住宿",
        "宾馆",
        "车票",
    ],
    "notice": [
        "通知",
        "会议通知",
        "培训通知",
        "录用通知",
        "邀请函",
        "公函",
        "派出公函",
        "来访",
        "受邀",
        "acceptance",
        "accepted",
        "conference",
        "讲座",
        "海报",
    ],
    "contract": [
        "合同",
        "协议",
        "contract",
        "订单",
        "采购订单",
        "服务协议",
    ],
    "detail_list": [
        "清单",
        "明细",
        "品名",
        "单价",
        "数量",
        "销货清单",
        "测试清单",
        "加工清单",
        "公务接待清单",
        "接待清单",
    ],
    "explanation": [
        "说明",
        "情况说明",
        "报销说明",
        "版面费说明",
    ],
}


def load_rules(path: Path = RULES_PATH) -> dict[str, Any]:
    return json.loads(path.read_text(encoding="utf-8"))


def clean_text(text: str) -> str:
    text = text.replace("\xa0", " ")
    text = re.sub(r"[ \t\r\f\v]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return "\n".join(line.strip() for line in text.splitlines() if line.strip())


def extract_text(path: Path) -> str:
    suffix = path.suffix.lower()
    if suffix == ".pdf":
        reader = PdfReader(str(path))
        pages = []
        for page in reader.pages:
            pages.append(page.extract_text() or "")
        return clean_text("\n".join(pages))
    if suffix == ".docx" and DocxDocument is not None:
        doc = DocxDocument(str(path))
        return clean_text("\n".join(p.text for p in doc.paragraphs))
    if suffix in {".txt", ".md", ".csv", ".json", ".xml", ".html", ".htm"}:
        return clean_text(path.read_text(encoding="utf-8", errors="ignore"))
    if suffix in {".png", ".jpg", ".jpeg", ".webp", ".tif", ".tiff", ".bmp"} and Image is not None and pytesseract is not None:
        try:
            return clean_text(pytesseract.image_to_string(Image.open(path), lang="chi_sim+eng"))
        except Exception:
            return clean_text(path.name)
    return clean_text(path.name)


def file_id(path: Path) -> str:
    safe = re.sub(r"[^a-zA-Z0-9\u4e00-\u9fff]+", "-", path.stem).strip("-")
    return safe[:60] or "document"


def score_keywords(text: str, keywords: list[str]) -> int:
    lower = text.lower()
    return sum(lower.count(k.lower()) for k in keywords)


def classify_document(path: Path, text: str) -> tuple[str, float]:
    haystack = f"{path.name}\n{text[:4000]}"
    scores = {
        doc_type: score_keywords(haystack, keywords)
        for doc_type, keywords in DOCUMENT_TYPE_KEYWORDS.items()
    }
    best_type, best_score = max(scores.items(), key=lambda item: item[1])
    if best_score == 0:
        return "other", 0.2
    total = sum(scores.values()) or 1
    return best_type, min(0.95, 0.45 + best_score / total)


def first_match(text: str, patterns: list[str], flags: int = 0) -> str:
    for pattern in patterns:
        match = re.search(pattern, text, flags)
        if not match:
            continue
        value = next((group for group in match.groups() if group), "")
        value = re.sub(r"\s+", " ", value).strip(" ：:，,")
        if value:
            return value
    return ""


def normalized_amount(value: str) -> float | None:
    try:
        return float(value.replace(",", ""))
    except ValueError:
        return None


def extract_fields(text: str, document_type: str = "") -> dict[str, Any]:
    fields: dict[str, Any] = {}
    amount_matches: list[str] = []
    amount_patterns = [
        r"(?:价税合计|金额|合计|小计|Amount|Total)[:：]?\s*(USD|CNY|RMB|¥|￥)?\s*([0-9]+(?:,[0-9]{3})*(?:\.[0-9]{1,2})?)\s*(元|USD|CNY|RMB)?",
        r"(?:¥|￥|USD|CNY|RMB)\s*([0-9]+(?:,[0-9]{3})*(?:\.[0-9]{1,2})?)",
        r"([0-9]+(?:,[0-9]{3})*(?:\.[0-9]{1,2})?)\s*(?:元|USD|CNY|RMB)",
    ]
    for pattern in amount_patterns:
        for match in re.findall(pattern, text, flags=re.I):
            if isinstance(match, tuple):
                amount_matches.append(next((part for part in match if re.search(r"\d", part)), ""))
                currency = next((part for part in match if part and re.search(r"USD|CNY|RMB|¥|￥|元", part, flags=re.I)), "")
                if currency and "currency" not in fields:
                    fields["currency"] = "CNY" if currency in {"¥", "￥", "元"} else currency.upper()
            else:
                amount_matches.append(match)
    normalized: list[float] = []
    if amount_matches:
        for match in amount_matches[:20]:
            value = normalized_amount(match)
            if value is not None and value <= 10_000_000:
                normalized.append(value)
        normalized = list(dict.fromkeys(normalized))[:10]
    if normalized:
        fields["amount_candidates"] = normalized
        fields["amount_total"] = max(normalized)
        fields["max_amount"] = fields["amount_total"]
    date_matches = re.findall(r"(20\d{2})[-./年](\d{1,2})[-./月](\d{1,2})", text)
    if date_matches:
        fields["date_candidates"] = [f"{int(y):04d}-{int(m):02d}-{int(d):02d}" for y, m, d in date_matches[:10]]
        fields["invoice_date"] = fields["date_candidates"][0]
    buyer_name = first_match(text, [r"(?:购买方名称|购买方|买方|付款方|客户名称)[:：]\s*([^\n]+)"])
    seller_name = first_match(text, [r"(?:销售方名称|销售方|卖方|收款方|供应商)[:：]\s*([^\n]+)"])
    item_name = first_match(text, [r"(?:项目名称|货物或应税劳务、服务名称|品名|商品名称|Item|Description)[:：]\s*([^\n]+)"], re.I)
    invoice_no = first_match(
        text,
        [
            r"(?:发票号码|发票号|票据号码|No\.?|Invoice\s*No\.?)[:：#]?\s*([A-Za-z0-9-]{6,})",
            r"\b([0-9]{16,24})\b",
        ],
        re.I,
    )
    if buyer_name:
        fields["buyer_name"] = buyer_name
    if seller_name:
        fields["seller_name"] = seller_name
    if item_name:
        fields["item_name"] = item_name
    if invoice_no:
        fields["invoice_no"] = invoice_no
        fields["invoice_no_detected"] = True
    if "上海大学" in text:
        fields["mentions_shu"] = True
    if document_type == "invoice":
        fields["invoice_kind"] = "foreign_invoice" if re.search(r"\b(invoice|receipt)\b", text, re.I) and "发票" not in text[:200] else "domestic_invoice"
        if fields["invoice_kind"] == "domestic_invoice":
            fields["title_valid"] = buyer_name == "上海大学" if buyer_name else fields.get("mentions_shu", False)
        else:
            fields["title_valid"] = None
    return fields


def import_documents(input_dir: Path) -> list[ImportedDocument]:
    if not input_dir.exists():
        raise FileNotFoundError(f"Input directory does not exist: {input_dir}")
    docs: list[ImportedDocument] = []
    for path in sorted(p for p in input_dir.rglob("*") if p.is_file()):
        if path.name.startswith("."):
            continue
        if "raw_emails" in path.parts or path.name == "email_metadata.json":
            continue
        text = extract_text(path)
        doc_type, confidence = classify_document(path, text)
        fields = extract_fields(text, doc_type)
        model_fields = extract_fields_with_doubao(path, text, doc_type)
        if model_fields.get("model_extract_error"):
            fields["model_extract_error"] = model_fields["model_extract_error"]
        else:
            for key, value in model_fields.items():
                fields.setdefault(key, value)
        docs.append(
            ImportedDocument(
                id=file_id(path),
                file_name=path.name,
                source_path=str(path.resolve()),
                document_type=doc_type,
                confidence=round(confidence, 3),
                text_preview=text[:500],
                extracted_fields=fields,
            )
        )
    return docs


def classify_case(rules: dict[str, Any], docs: list[ImportedDocument], user_note: str) -> tuple[dict[str, Any], float, dict[str, int]]:
    combined = "\n".join([user_note] + [d.file_name + "\n" + d.text_preview for d in docs])
    scores = {
        rule["id"]: score_keywords(combined, rule.get("keywords", []))
        for rule in rules["supported_types"]
    }
    best_id, best_score = max(scores.items(), key=lambda item: item[1])
    if best_score == 0:
        # Default to daily supplies because it is the least workflow-specific form.
        best_id = "office_supplies"
        best_score = 1
    best_rule = next(rule for rule in rules["supported_types"] if rule["id"] == best_id)
    total = sum(scores.values()) or best_score
    confidence = min(0.95, max(0.35, best_score / total))
    return best_rule, round(confidence, 3), scores


def evidence_for_query(index: dict[str, Any], query: str) -> SourceEvidence | None:
    hits = bm25_scores(index, query, 1)
    if not hits:
        return None
    score, chunk = hits[0]
    return SourceEvidence(
        title=chunk["title"],
        url=chunk["url"],
        source_type=chunk["source_type"],
        snippet=snippet(chunk["text"], query, 260),
        score=round(float(score), 3),
    )


def doc_matches_item(doc: ImportedDocument, item: dict[str, Any]) -> bool:
    if doc.document_type in item.get("accepted_document_types", []):
        keywords = item.get("keywords") or []
        if not keywords:
            return True
        haystack = f"{doc.file_name}\n{doc.text_preview}"
        return score_keywords(haystack, keywords) > 0
    return False


def build_checklist(rule: dict[str, Any], docs: list[ImportedDocument], index: dict[str, Any]) -> list[ChecklistItem]:
    checklist = []
    for item in rule["checklist"]:
        matched = [doc.id for doc in docs if doc_matches_item(doc, item)]
        if matched:
            status = "ready"
        elif item["required"]:
            status = "missing"
        else:
            status = "optional"
        query = f"{rule['name']} {item['name']} {item.get('why', '')}"
        evidence = evidence_for_query(index, query)
        checklist.append(
            ChecklistItem(
                id=item["id"],
                name=item["name"],
                required=bool(item["required"]),
                status=status,
                matched_documents=matched,
                why=item.get("why", ""),
                evidence=evidence,
            )
        )
    return checklist


def build_policy_evidence(rule: dict[str, Any], index: dict[str, Any]) -> list[SourceEvidence]:
    evidence = []
    seen = set()
    for query in rule.get("policy_queries", []):
        item = evidence_for_query(index, query)
        if item and item.url not in seen:
            evidence.append(item)
            seen.add(item.url)
    return evidence


def infer_purpose(user_note: str, rule: dict[str, Any], docs: list[ImportedDocument]) -> str:
    if user_note:
        return user_note.strip("。.")
    if rule["id"] == "domestic_travel":
        return "开展出差、会议、调研或学术交流"
    if rule["id"] == "publication_fee":
        return "论文发表、审稿或出版"
    if rule["id"] == "office_supplies":
        return "教学科研或日常办公需要"
    if rule["id"] == "materials_fee":
        return "教学科研实验材料采购"
    if rule["id"] == "testing_processing_fee":
        return "科研测试、化验或加工服务"
    if rule["id"] == "official_reception":
        return "公务接待"
    if rule["id"] == "information_service_fee":
        return "信息服务、数据服务或软件服务"
    return "相关业务"


def reason_draft(rule: dict[str, Any], user_note: str, docs: list[ImportedDocument]) -> str:
    purpose = infer_purpose(user_note, rule, docs)
    return rule["reason_template"].format(purpose=purpose)


def risk_flags(docs: list[ImportedDocument], checklist: list[ChecklistItem]) -> list[str]:
    flags = []
    missing = [item.name for item in checklist if item.status == "missing"]
    if missing:
        flags.append("缺少必需材料：" + "、".join(missing))
    invoice_docs = [d for d in docs if d.document_type == "invoice"]
    domestic_invoices = [d for d in invoice_docs if d.extracted_fields.get("invoice_kind") != "foreign_invoice"]
    invalid_title = [d.file_name for d in domestic_invoices if d.extracted_fields.get("title_valid") is False]
    if invalid_title:
        flags.append("发票抬头可能不是上海大学：" + "、".join(invalid_title) + "。")
    unknown_title = [d.file_name for d in domestic_invoices if "title_valid" not in d.extracted_fields]
    if unknown_title:
        flags.append("部分发票未能识别购买方抬头，请人工核对：" + "、".join(unknown_title) + "。")
    missing_invoice_fields = []
    for doc in invoice_docs:
        missing_fields = []
        if not doc.extracted_fields.get("invoice_no"):
            missing_fields.append("发票号")
        if not doc.extracted_fields.get("invoice_date"):
            missing_fields.append("日期")
        if not doc.extracted_fields.get("amount_total"):
            missing_fields.append("金额")
        if missing_fields:
            missing_invoice_fields.append(f"{doc.file_name} 缺 {','.join(missing_fields)}")
    if missing_invoice_fields:
        flags.append("发票关键字段识别不完整：" + "；".join(missing_invoice_fields) + "。")
    if not any(d.document_type == "payment_record" for d in docs):
        flags.append("未识别到支付记录，提交前通常需要补充公务卡、银行卡、POS 小票或银行记录。")
    if any("纸质" in d.text_preview or "出租车" in d.text_preview for d in docs):
        flags.append("可能包含纸质票据，线上整理后仍需按上海大学财务系统要求线下投递。")
    return flags


def slugify(value: str) -> str:
    slug = re.sub(r"[^a-zA-Z0-9\u4e00-\u9fff]+", "-", value).strip("-")
    return slug[:80] or "reimbursement-package"


DOC_TYPE_LABELS = {
    "invoice": "发票",
    "payment_record": "支付记录",
    "detail_list": "明细清单",
    "approval": "审批/申请",
    "settlement": "结算单",
    "itinerary": "行程/票据",
    "notice": "通知/证明",
    "contract": "合同",
    "other": "其他材料",
}

CHECK_STATUS_LABELS = {
    "ready": "已就绪",
    "missing": "缺失",
    "optional": "可选/按需",
}


def md_escape(value: Any) -> str:
    text = str(value if value is not None else "")
    text = text.replace("|", "\\|")
    text = re.sub(r"\s+", " ", text).strip()
    return text


def format_amount(fields: dict[str, Any]) -> str:
    amount = fields.get("amount_total") or fields.get("max_amount")
    if amount in ("", None):
        return "-"
    currency = fields.get("currency") or "CNY"
    try:
        amount_value = float(amount)
        amount_text = f"{amount_value:,.2f}".rstrip("0").rstrip(".")
    except (TypeError, ValueError):
        amount_text = str(amount)
    return f"{currency} {amount_text}"


def title_check_text(fields: dict[str, Any]) -> str:
    if fields.get("title_valid") is True:
        return "抬头有效"
    if fields.get("title_valid") is False:
        return "抬头待处理"
    if fields.get("invoice_kind") == "foreign_invoice":
        return "外文票据"
    return "需人工核对"


def document_review_points(doc: ImportedDocument) -> str:
    fields = doc.extracted_fields
    points: list[str] = []
    if doc.document_type == "invoice":
        points.append(title_check_text(fields))
        if not fields.get("invoice_no"):
            points.append("核对发票号")
        if not fields.get("invoice_date"):
            points.append("核对开票日期")
        if not fields.get("amount_total"):
            points.append("核对金额")
    if doc.document_type == "payment_record":
        points.append("核对付款人与金额")
    if doc.document_type in {"detail_list", "contract"}:
        points.append("核对品名、数量、单价")
    if not points:
        points.append("核对原件内容")
    return "；".join(dict.fromkeys(points))


def source_label(evidence: SourceEvidence) -> str:
    title = clean_source_title(evidence.title)
    if evidence.url.startswith(("http://", "https://")):
        return f"[{md_escape(title)}]({evidence.url})"
    return f"{md_escape(title)}（本地资料）"


def evidence_key(evidence: SourceEvidence | None) -> tuple[str, str]:
    if evidence is None:
        return ("", "")
    return (evidence.title, evidence.url)


def evidence_numbers(package: ReimbursementPackage) -> dict[tuple[str, str], int]:
    ordered: list[SourceEvidence] = []
    seen: set[tuple[str, str]] = set()
    for item in package.checklist:
        key = evidence_key(item.evidence)
        if item.evidence and key not in seen:
            seen.add(key)
            ordered.append(item.evidence)
    for evidence in package.policy_evidence:
        key = evidence_key(evidence)
        if key not in seen:
            seen.add(key)
            ordered.append(evidence)
    return {evidence_key(evidence): idx for idx, evidence in enumerate(ordered, 1)}


def evidence_ref(evidence: SourceEvidence | None, refs: dict[tuple[str, str], int]) -> str:
    if evidence is None:
        return "-"
    number = refs.get(evidence_key(evidence))
    return f"依据 {number}" if number else "-"


def estimated_total_amount(docs: list[ImportedDocument]) -> float:
    preferred = [doc for doc in docs if doc.document_type == "invoice"]
    if not preferred:
        preferred = [doc for doc in docs if doc.document_type == "payment_record"]
    if not preferred:
        preferred = docs
    total = 0.0
    for doc in preferred:
        amount = doc.extracted_fields.get("amount_total")
        if isinstance(amount, (int, float)):
            total += float(amount)
    return total


def clean_source_title(title: str) -> str:
    return re.sub(r"^[0-9a-f]{8,16}-", "", title).strip() or title


def markdown_report(package: ReimbursementPackage) -> str:
    missing_required = [item for item in package.checklist if item.required and item.status == "missing"]
    ready_items = [item for item in package.checklist if item.status == "ready"]
    optional_items = [item for item in package.checklist if item.status == "optional"]
    total_amount = estimated_total_amount(package.documents)
    refs = evidence_numbers(package)
    created_at = package.created_at.replace("T", " ").split(".")[0]

    lines = [
        f"# {package.title}",
        "",
        "## 一页结论",
        "",
        f"- 报销类型：{package.reimbursement_type_name}",
        f"- 推荐入口：{package.school_form_hint}",
        f"- 当前状态：{'待补材料' if missing_required else '材料可进入本人复核'}",
        f"- 已识别材料：{len(package.documents)} 份",
        f"- 必需材料缺失：{len(missing_required)} 项",
        f"- 识别金额合计：{'CNY ' + format(total_amount, ',.2f') if total_amount else '未汇总'}",
        f"- 分类置信度：{package.confidence}",
        f"- 生成时间：{created_at}",
        f"- 用户备注：{package.user_note or '无'}",
        "",
        "## 报销事由草稿",
        "",
        package.reason_draft,
        "",
        "## 下一步处理",
        "",
    ]

    if missing_required:
        lines.append("先补齐以下必需材料，再进入学校财务系统提交：")
        lines.append("")
        for item in missing_required:
            lines.append(f"- {item.name}：{item.why}（{evidence_ref(item.evidence, refs)}）")
    else:
        lines.extend(
            [
                "- 材料已满足当前规则检查，可以进入本人复核。",
                "- 提交前仍需核对经费项目、收款账号、发票真伪、附件完整性和学校系统预审提示。",
            ]
        )
    if optional_items:
        lines.append("")
        lines.append("按业务情况确认的材料：")
        for item in optional_items:
            lines.append(f"- {item.name}：{item.why}")

    lines.extend(
        [
            "",
            "## 材料识别",
            "",
            "| 文件 | 类型 | 金额 | 日期/号码 | 复核点 |",
            "| --- | --- | ---: | --- | --- |",
        ]
    )
    for doc in package.documents:
        fields = doc.extracted_fields
        date_no = []
        if fields.get("invoice_date"):
            date_no.append(f"日期 {fields['invoice_date']}")
        elif fields.get("transaction_date"):
            date_no.append(f"日期 {fields['transaction_date']}")
        if fields.get("invoice_no"):
            date_no.append(f"发票号 {fields['invoice_no']}")
        if fields.get("buyer_name"):
            date_no.append(f"购买方 {fields['buyer_name']}")
        lines.append(
            "| "
            + " | ".join(
                [
                    md_escape(doc.file_name),
                    md_escape(DOC_TYPE_LABELS.get(doc.document_type, doc.document_type)),
                    md_escape(format_amount(fields)),
                    md_escape("；".join(date_no) or "-"),
                    md_escape(document_review_points(doc)),
                ]
            )
            + " |"
        )

    lines.extend(["", "## 完整性检查", "", "| 检查项 | 状态 | 匹配材料 | 说明 | 依据 |", "| --- | --- | --- | --- | --- |"])
    for item in package.checklist:
        matched = "、".join(item.matched_documents) if item.matched_documents else "-"
        lines.append(
            "| "
            + " | ".join(
                [
                    md_escape(item.name),
                    md_escape(CHECK_STATUS_LABELS.get(item.status, item.status)),
                    md_escape(matched),
                    md_escape(item.why),
                    md_escape(evidence_ref(item.evidence, refs)),
                ]
            )
            + " |"
        )

    lines.extend(["", "## 发票与金额核对", ""])
    invoice_docs = [doc for doc in package.documents if doc.document_type == "invoice"]
    if invoice_docs:
        for doc in invoice_docs:
            fields = doc.extracted_fields
            checks = [
                f"文件：{doc.file_name}",
                f"抬头：{fields.get('buyer_name') or '待核对'}",
                f"销售方：{fields.get('seller_name') or '待核对'}",
                f"金额：{format_amount(fields)}",
                f"日期：{fields.get('invoice_date') or '待核对'}",
                f"发票号：{fields.get('invoice_no') or '待核对'}",
                f"结论：{title_check_text(fields)}",
            ]
            lines.append("- " + "；".join(checks))
    else:
        lines.append("- 未识别到发票材料，请人工确认是否需要补充发票或财政票据。")

    lines.extend(["", "## 风险与人工确认", ""])
    if package.risk_flags:
        lines.extend(f"- {flag}" for flag in package.risk_flags)
    else:
        lines.append("- 暂未发现高优先级风险。")
    lines.extend(
        [
            "- 提交前必须由本人确认票据真实合法、业务真实发生、附件与报销事项一致。",
            "- 经费项目、预算额度、支付账号、冲借款信息和审批流必须在学校系统中人工核对。",
            "",
            "## 政策依据",
            "",
        ]
    )

    if refs:
        evidence_by_ref: dict[int, SourceEvidence] = {}
        for item in package.checklist:
            if item.evidence:
                number = refs.get(evidence_key(item.evidence))
                if number:
                    evidence_by_ref[number] = item.evidence
        for evidence in package.policy_evidence:
            number = refs.get(evidence_key(evidence))
            if number:
                evidence_by_ref[number] = evidence
        for number in sorted(evidence_by_ref):
            evidence = evidence_by_ref[number]
            clean_snippet = re.sub(r"\s+", " ", evidence.snippet).strip()
            lines.append(f"### 依据 {number}")
            lines.append("")
            lines.append(f"- 来源：{source_label(evidence)}")
            lines.append(f"- 类型：{evidence.source_type}")
            lines.append(f"- 相关片段：{clean_snippet[:420]}{'...' if len(clean_snippet) > 420 else ''}")
            lines.append("")
    else:
        lines.append("- 暂无可追溯政策依据。")

    lines.extend(
        [
            "## 提交边界",
            "",
            "- 本文件是本地整理结果，不代表财务处最终审核意见。",
            "- 登录、验证码、经费项目、收款账号、提交预审和提交审批均需本人确认。",
            "- Agent 只能辅助整理材料和填写草稿，不能代替本人作真实性承诺或最终提交。",
        ]
    )
    return "\n".join(lines) + "\n"


def playwright_plan(package: ReimbursementPackage) -> dict[str, Any]:
    reason_short = package.user_note.strip() or package.reason_draft.strip()
    note_short = package.reason_draft.strip()
    ready_documents = [
        {
            "id": doc.id,
            "file_name": doc.file_name,
            "document_type": doc.document_type,
            "source_path": doc.source_path,
        }
        for doc in package.documents
    ]
    missing_required = [
        {
            "id": item.id,
            "name": item.name,
            "why": item.why,
            "evidence_url": item.evidence.url if item.evidence else "",
        }
        for item in package.checklist
        if item.status == "missing"
    ]
    return {
        "package_id": package.id,
        "title": package.title,
        "status": "ready_for_human_review" if not missing_required else "needs_materials_before_submission",
        "school_system": {
            "recommended_entry": package.school_form_hint,
            "reimbursement_type": package.reimbursement_type_name,
        },
        "draft_fields": {
            "reason": package.reason_draft,
            "reason_short": reason_short[:30],
            "note_short": note_short[:60],
            "user_note": package.user_note,
        },
        "upload_documents": ready_documents,
        "missing_required": missing_required,
        "must_pause_for_confirmation": [
            "登录、验证码和二次认证",
            "选择经费项目",
            "确认银行卡或对公账号",
            "冲借款",
            "提交预审",
            "提交审批",
            "任何真实性、合规性或承诺类勾选",
        ],
        "agent_boundary": "Playwright 只能辅助录入和上传。提交预审、提交审批以及真实性承诺必须由本人确认。",
    }


def build_package(input_dir: Path, user_note: str, title: str | None = None) -> ReimbursementPackage:
    rules = load_rules()
    index = load_index(INDEX_PATH)
    docs = import_documents(input_dir)
    rule, confidence, _scores = classify_case(rules, docs, user_note)
    checklist = build_checklist(rule, docs, index)
    policy = build_policy_evidence(rule, index)
    package_title = title or f"{rule['name']}报销包"
    return ReimbursementPackage(
        id=slugify(package_title),
        title=package_title,
        created_at=datetime.now(timezone.utc).isoformat(),
        input_dir=str(input_dir.resolve()),
        reimbursement_type=rule["id"],
        reimbursement_type_name=rule["name"],
        school_form_hint=rule["school_form_hint"],
        confidence=confidence,
        user_note=user_note,
        reason_draft=reason_draft(rule, user_note, docs),
        documents=docs,
        checklist=checklist,
        risk_flags=risk_flags(docs, checklist),
        policy_evidence=policy,
    )


def export_package(package: ReimbursementPackage, out_dir: Path) -> None:
    out_dir.mkdir(parents=True, exist_ok=True)
    files_dir = out_dir / "files"
    if files_dir.exists():
        shutil.rmtree(files_dir)
    files_dir.mkdir(exist_ok=True)
    for doc in package.documents:
        src = Path(doc.source_path)
        if src.exists():
            shutil.copy2(src, files_dir / src.name)
    (out_dir / "package.json").write_text(json.dumps(asdict(package), ensure_ascii=False, indent=2), encoding="utf-8")
    (out_dir / "package.md").write_text(markdown_report(package), encoding="utf-8")
    (out_dir / "playwright_plan.json").write_text(
        json.dumps(playwright_plan(package), ensure_ascii=False, indent=2),
        encoding="utf-8",
    )


def main() -> None:
    parser = argparse.ArgumentParser(description="Build a reimbursement package from local files.")
    parser.add_argument("input_dir", type=Path)
    parser.add_argument("--note", default="", help="User's free-form reimbursement context.")
    parser.add_argument("--title", default="")
    parser.add_argument("--out", type=Path, default=None)
    args = parser.parse_args()

    package = build_package(args.input_dir, args.note, args.title or None)
    out_dir = args.out or OUTPUT_ROOT / package.id
    export_package(package, out_dir)
    summary = {
        "package": str(out_dir),
        "type": package.reimbursement_type_name,
        "confidence": package.confidence,
        "documents": len(package.documents),
        "missing_required": [item.name for item in package.checklist if item.status == "missing"],
        "reason_draft": package.reason_draft,
    }
    print(json.dumps(summary, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
